"""
PPT Agent — builds the Perfect Store deck (python-pptx).

One crisp slide per segment: stat strip, a data-driven (differentiated) radar,
a 2-line shopper snapshot, the top-3 actions and hero SKUs. Plus a title slide,
a portfolio-overview slide and an A/B/C/D priority snapshot.

Entry point:  generate_pptx(df, labels, output_dir, radar_paths=..., priority_narrative=...)
"""

from __future__ import annotations

import logging
import os
from datetime import datetime

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

logger = logging.getLogger("perfect_store.ppt")

_PALETTE = ["004B87", "009CDE", "E4002B", "FFB81C",
            "41B6E6", "6CC24A", "FF6900", "7B2D8B"]
_INK = RGBColor.from_string("1F2937")
_MUTE = RGBColor.from_string("6B7280")
_WHITE = RGBColor.from_string("FFFFFF")
_PRIORITY_ORDER = ["A+", "A", "B+", "B", "C+", "C", "D+", "D"]
_SEC_COLS = ["sec_a_hhs", "sec_b_hhs", "sec_c_hhs", "sec_d_hhs", "sec_e_hhs"]


# ── low-level helpers ──────────────────────────────────────────────────────

def _text(slide, x, y, w, h, runs, *, size=14, bold=False, color=_INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.05):
    """runs: str, or list of (text, {overrides}) tuples, or list of paragraphs
    where a paragraph is itself such a list."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        for txt, ov in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(ov.get("size", size))
            r.font.bold = ov.get("bold", bold)
            r.font.name = "Calibri"
            r.font.color.rgb = ov.get("color", color)
    return tb


def _rect(slide, x, y, w, h, hexcolor):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(hexcolor)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rupees(v) -> str:
    try:
        return f"₹{float(v):,.0f}"
    except (TypeError, ValueError):
        return "—"


def _mean(sub: pd.DataFrame, col: str):
    if col not in sub.columns:
        return None
    s = pd.to_numeric(sub[col], errors="coerce")
    return s.mean() if s.notna().any() else None


def _dominant_sec(sub: pd.DataFrame) -> str:
    present = [c for c in _SEC_COLS if c in sub.columns]
    if not present:
        return ""
    means = {c: pd.to_numeric(sub[c], errors="coerce").mean() for c in present}
    best = max(means, key=lambda c: (means[c] if means[c] == means[c] else -1))
    return best.split("_")[1].upper()


# ── card fallback ─────────────────────────────────────────────────────────

def _fallback_card(info: dict) -> dict:
    return {
        "headline": info.get("label", ""),
        "shopper_snapshot": (info.get("description") or "").strip()[:240],
        "mission": info.get("occasion", ""),
        "growth_potential": "Medium",
        "dominant_sec": "",
        "actions": [{
            "lever": "Execution",
            "text": (info.get("action") or "Review segment priorities").strip()[:150],
            "kpi": "",
        }],
        "hero_skus": [],
    }


# ── slides ────────────────────────────────────────────────────────────────

def _title_slide(prs, df, labels):
    s = _blank(prs)
    _rect(s, 0, 0, prs.slide_width / 914400, 1.6, _PALETTE[0])
    _text(s, 0.7, 0.35, 12, 1.0, "Perfect Store", size=40, bold=True, color=_WHITE)
    _text(s, 0.72, 1.15, 12, 0.6, "Segmentation & Activation Deck",
          size=18, color=RGBColor.from_string("D6E6F2"))

    n_out = len(df)
    n_seg = df["cluster"].nunique()
    tot_vpo = pd.to_numeric(df.get("VPO"), errors="coerce").sum()
    _text(s, 0.72, 2.5, 12, 0.5,
          [[(f"{n_out:,}", {"bold": True, "size": 16}), ("  outlets", {}),
            ("      ", {}),
            (f"{n_seg}", {"bold": True, "size": 16}), ("  segments", {}),
            ("      ", {}),
            (_rupees(tot_vpo), {"bold": True, "size": 16}),
            ("  total monthly VPO", {})]],
          size=14, color=_INK)
    _text(s, 0.72, 6.7, 12, 0.4,
          f"Generated {datetime.now():%d %b %Y  %H:%M}", size=11, color=_MUTE)


def _overview_slide(prs, df, labels, output_dir):
    s = _blank(prs)
    _text(s, 0.5, 0.3, 12.3, 0.6, "Portfolio Overview", size=26, bold=True,
          color=_PALETTE_INK)

    cids = sorted(df["cluster"].unique())
    headers = ["Seg", "Label", "Channel", "Outlets", "% Univ", "Avg VPO", "Avg SKUs", "Growth"]
    rows = []
    for cid in cids:
        sub = df[df["cluster"] == cid]
        info = labels.get(int(cid), {})
        card = info.get("card") or {}
        rows.append([
            str(int(cid)),
            (info.get("label", f"Segment {cid}"))[:34],
            info.get("channel", "") or "—",
            f"{len(sub):,}",
            f"{100 * len(sub) / len(df):.1f}%",
            _rupees(_mean(sub, "VPO")),
            f"{_mean(sub, 'AVG_SKU'):.1f}" if _mean(sub, "AVG_SKU") is not None else "—",
            card.get("growth_potential", "") or "—",
        ])

    _table(s, 0.5, 1.1, 8.4, headers, rows,
           col_widths=[0.5, 3.0, 1.0, 1.0, 0.9, 1.2, 1.0, 1.0])

    bar = os.path.join(output_dir, "charts", "cluster_bar_overview.png")
    if os.path.exists(bar):
        s.shapes.add_picture(bar, Inches(9.2), Inches(1.2), width=Inches(3.8))


def _priority_slide(prs, df):
    if "priority" not in df.columns:
        return
    s = _blank(prs)
    _text(s, 0.5, 0.3, 12.3, 0.6, "Store Prioritization — A / B / C / D",
          size=26, bold=True, color=_PALETTE_INK)

    headers = ["Tier", "Outlets", "% Univ", "Avg Actual VPO", "Avg Potential VPO", "Avg Gap"]
    rows = []
    for tier in _PRIORITY_ORDER:
        sub = df[df["priority"] == tier]
        if not len(sub):
            continue
        rows.append([
            tier,
            f"{len(sub):,}",
            f"{100 * len(sub) / len(df):.1f}%",
            _rupees(_mean(sub, "VPO")),
            _rupees(_mean(sub, "predicted_potential_vpo")),
            _rupees(_mean(sub, "opportunity_gap")),
        ])
    if not rows:
        return
    _table(s, 0.5, 1.2, 9.0, headers, rows,
           col_widths=[1.0, 1.4, 1.2, 2.1, 2.2, 1.6])
    _text(s, 0.5, 6.7, 12, 0.4,
          "A/B = top 40% by VPO · C = next 30% · D = bottom 30%. "
          "“+” tier = below modelled 75th-percentile potential (highest upside).",
          size=10, color=_MUTE)


def _segment_slide(prs, df, cid, info, radar_path):
    s = _blank(prs)
    colour = _PALETTE[int(cid) % len(_PALETTE)]
    _rect(s, 0, 0, prs.slide_width / 914400, 0.18, colour)

    sub = df[df["cluster"] == cid]
    card = info.get("card") or _fallback_card(info)
    label = info.get("label", f"Segment {cid}")
    pct = 100 * len(sub) / len(df)

    _text(s, 0.5, 0.32, 12.4, 0.7, f"Segment {int(cid)} — {label}",
          size=27, bold=True, color=RGBColor.from_string(colour))
    sub_bits = " · ".join(
        b for b in [info.get("channel", ""), info.get("occasion", ""),
                    f"{len(sub):,} outlets ({pct:.1f}% of universe)"] if b)
    _text(s, 0.5, 1.02, 12.4, 0.4, sub_bits, size=13, color=_MUTE)

    if card.get("headline"):
        _text(s, 0.5, 1.42, 12.4, 0.45,
              [[(card["headline"], {"bold": True, "size": 14,
                                    "color": RGBColor.from_string(colour)})]])

    # ── left: stat strip + radar ──
    sec = card.get("dominant_sec") or _dominant_sec(sub)
    stat_rows = [
        ("Avg monthly VPO", _rupees(_mean(sub, "VPO"))),
        ("Avg total revenue", _rupees(_mean(sub, "TOTAL_REVENUE"))),
        ("Avg SKUs stocked", f"{_mean(sub, 'AVG_SKU'):.1f}" if _mean(sub, "AVG_SKU") is not None else "—"),
        ("Avg active months", f"{_mean(sub, 'ACTIVE_MONTHS'):.1f}" if _mean(sub, "ACTIVE_MONTHS") is not None else "—"),
        ("Growth potential", card.get("growth_potential", "—") or "—"),
        ("Dominant SEC", sec or "—"),
    ]
    _table(s, 0.5, 2.05, 5.4, ["Metric", "Value"], stat_rows,
           col_widths=[3.3, 2.1], header=False, compact=True)

    if radar_path and os.path.exists(radar_path):
        s.shapes.add_picture(radar_path, Inches(0.35), Inches(4.2),
                             height=Inches(2.85))

    # ── right: snapshot / mission / actions / hero SKUs ──
    rx, rw = 6.35, 6.5
    y = 2.05
    _text(s, rx, y, rw, 0.35, "WHO SHOPS HERE", size=11, bold=True, color=_MUTE)
    _text(s, rx, y + 0.32, rw, 1.1, card.get("shopper_snapshot", "") or "—",
          size=12.5)
    y += 1.55
    if card.get("mission"):
        _text(s, rx, y, rw, 0.4,
              [[("PRIMARY MISSION   ", {"size": 11, "bold": True, "color": _MUTE}),
                (card["mission"], {"size": 12.5})]])
        y += 0.5

    _text(s, rx, y, rw, 0.35, "TOP ACTIONS", size=11, bold=True, color=_MUTE)
    y += 0.34
    for a in (card.get("actions") or [])[:3]:
        if isinstance(a, dict):
            lever = a.get("lever", "").strip()
            txt = a.get("text", "").strip()
            kpi = a.get("kpi", "").strip()
        else:
            lever, txt, kpi = "", str(a), ""
        runs = [("▸ ", {"color": RGBColor.from_string(colour), "bold": True})]
        if lever:
            runs.append((f"{lever}: ", {"bold": True}))
        runs.append((txt, {}))
        if kpi:
            runs.append((f"  (KPI: {kpi})", {"color": _MUTE, "size": 10.5}))
        _text(s, rx, y, rw, 0.5, [runs], size=12)
        y += 0.52

    heroes = card.get("hero_skus") or []
    if heroes:
        y += 0.05
        _text(s, rx, y, rw, 0.7,
              [[("HERO SKUs   ", {"size": 11, "bold": True, "color": _MUTE}),
                (", ".join(str(h) for h in heroes), {"size": 12})]])

    _text(s, 6.35, 7.15, 6.5, 0.3, f"Perfect Store  ·  {datetime.now():%d %b %Y}",
          size=9, color=_MUTE, align=PP_ALIGN.RIGHT)


# ── table helper ─────────────────────────────────────────────────────────

def _table(slide, x, y, total_w, headers, rows, *, col_widths=None,
           header=True, compact=False):
    n_col = len(headers)
    n_row = len(rows) + (1 if header else 0)
    row_h = 0.28 if compact else 0.34
    gt = slide.shapes.add_table(n_row, n_col, Inches(x), Inches(y),
                                Inches(total_w), Inches(row_h * n_row)).table
    gt.first_row = header
    gt.horz_banding = True
    if col_widths:
        for i, w in enumerate(col_widths[:n_col]):
            gt.columns[i].width = Inches(w)
    r0 = 0
    if header:
        for j, htxt in enumerate(headers):
            c = gt.cell(0, j)
            c.text = str(htxt)
            _style_cell(c, bold=True, size=10, color=_WHITE,
                        fill="004B87")
        r0 = 1
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = gt.cell(r0 + i, j)
            c.text = str(val)
            _style_cell(c, bold=(j == 0 and not header),
                        size=10 if not compact else 10.5,
                        fill="F3F4F6" if (i % 2 == 0) else "FFFFFF")
    return gt


def _style_cell(cell, *, bold=False, size=10, color=_INK, fill=None):
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Calibri"
            r.font.color.rgb = color


_PALETTE_INK = RGBColor.from_string("1F2937")


# ── entry point ──────────────────────────────────────────────────────────

def generate_pptx(df: pd.DataFrame, labels: dict, output_dir: str,
                  radar_paths: dict | None = None,
                  priority_narrative: str = "") -> str:
    radar_paths = radar_paths or {}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _title_slide(prs, df, labels)
    _overview_slide(prs, df, labels, output_dir)
    _priority_slide(prs, df)

    for cid in sorted(df["cluster"].unique()):
        info = labels.get(int(cid), {})
        _segment_slide(prs, df, int(cid), info, radar_paths.get(int(cid)))

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(output_dir, f"perfect_store_deck_{ts}.pptx")
    prs.save(path)
    logger.info(f"Wrote PPTX deck: {path}  ({len(prs.slides)} slides)")
    return path
