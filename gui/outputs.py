"""Discover and render the deliverables a pipeline run drops in
``outputs/<timestamp>/`` — PPTX deck, Excel workbooks, segment CSVs, PNG charts.
"""

from __future__ import annotations

import base64
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st

from gui import ui

_PRIORITY_ORDER = ["A+", "A", "B+", "B", "C+", "C", "D+", "D"]


def list_runs(project_root: str) -> list[Path]:
    """Newest-first list of run folders under outputs/."""
    out = Path(project_root) / "outputs"
    if not out.is_dir():
        return []
    runs = [p for p in out.iterdir() if p.is_dir() and p.name[0].isdigit()]
    return sorted(runs, reverse=True)


def _arrow_safe(df: pd.DataFrame) -> pd.DataFrame:
    obj_cols = df.select_dtypes(include="object").columns
    if len(obj_cols):
        df = df.astype({c: "string" for c in obj_cols})
    return df


def _fmt_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}"
        n /= 1024
    return f"{n:.0f} TB"


def run_label(run_dir: Path) -> str:
    name = run_dir.name
    try:
        return datetime.strptime(name, "%Y%m%d_%H%M%S").strftime("%d %b %Y  %H:%M:%S")
    except ValueError:
        return name


def scan_artifacts(run_dir: Path) -> dict:
    run_dir = Path(run_dir)
    return {
        "pptx": sorted(run_dir.glob("*.pptx")),
        "pdf": sorted(run_dir.glob("*.pdf")),
        "excel": sorted(run_dir.glob("*.xlsx")),
        "csv": sorted(run_dir.glob("*.csv")),
        "charts": sorted((run_dir / "charts").glob("*.png"))
        if (run_dir / "charts").is_dir() else [],
    }


# ── download strip (artifact cards) ────────────────────────────────────────

_ICON = {"pptx": "🖥", "pdf": "📄", "xlsx": "📗", "csv": "📑"}
_MIME = {
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


def _slide_count(path: Path) -> str:
    try:
        from pptx import Presentation
        return f"{len(list(Presentation(str(path)).slides))} slides"
    except Exception:  # noqa: BLE001
        return ""


def _download_strip(art: dict) -> None:
    items: list[Path] = (art["pptx"] + art["pdf"] + art["excel"] + art["csv"])
    if not items:
        return
    st.markdown("###### Downloads")
    cols = st.columns(4)
    for i, p in enumerate(items):
        ext = p.suffix.lstrip(".").lower()
        badge = _slide_count(p) if ext == "pptx" else ""
        with cols[i % 4]:
            ui.art_card_header(_ICON.get(ext, "📁"), p.name,
                               _fmt_size(p.stat().st_size), badge)
            st.download_button("Download", p.read_bytes(), file_name=p.name,
                               mime=_MIME.get(ext, "application/octet-stream"),
                               key=f"dl_{p}", width="stretch")


# ── individual renderers ─────────────────────────────────────────────────────

def _render_pdf(path: Path) -> None:
    b64 = base64.b64encode(path.read_bytes()).decode()
    st.markdown(
        f'<iframe src="data:application/pdf;base64,{b64}" width="100%" '
        f'height="820" style="border:1px solid var(--ps-border);border-radius:8px">'
        f'</iframe>', unsafe_allow_html=True)
    st.caption("Use the download above if the preview is blank "
               "(some browsers block inline PDFs).")


def _render_pptx(path: Path) -> None:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = list(prs.slides)
    except Exception as exc:  # noqa: BLE001
        ui.callout("warn", "Deck saved", f"Slide list unavailable: {exc}")
        return
    st.caption(f"{len(slides)} slides · {_fmt_size(path.stat().st_size)}")
    titles = []
    for i, sl in enumerate(slides, 1):
        t = ""
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().splitlines()[0]
                break
        titles.append(f"{i}. {t or '(untitled)'}")
    st.markdown("\n".join(f"- {x}" for x in titles))
    ui.callout("info", "Open in a presentation app",
               "Download above, then open in PowerPoint / Google Slides / Keynote "
               "— browsers can't render .pptx inline.")


def _render_excel(path: Path) -> None:
    try:
        xl = pd.ExcelFile(path)
    except Exception as exc:  # noqa: BLE001
        ui.callout("error", f"Could not open {path.name}", str(exc))
        return
    sheet = st.selectbox("Sheet", xl.sheet_names, key=f"sheet_{path}")
    try:
        df = xl.parse(sheet)
        st.dataframe(_arrow_safe(df), width="stretch", height=460, hide_index=True)
        st.caption(f"{len(df):,} rows × {len(df.columns)} columns — "
                   "embedded charts are in the downloaded file.")
    except Exception as exc:  # noqa: BLE001
        ui.callout("error", f"Could not read sheet '{sheet}'", str(exc))


def _render_csv(path: Path) -> None:
    try:
        df = pd.read_csv(path)
    except Exception as exc:  # noqa: BLE001
        ui.callout("error", f"Could not read {path.name}", str(exc))
        return
    st.dataframe(_arrow_safe(df), width="stretch", height=420, hide_index=True)
    st.caption(f"{len(df):,} rows × {len(df.columns)} columns")


# ── top-level viewer ────────────────────────────────────────────────────────

def render_run(run_dir: Path) -> None:
    run_dir = Path(run_dir)
    art = scan_artifacts(run_dir)
    if not sum(len(v) for v in art.values()):
        ui.callout("empty", "No output files yet",
                   "This run folder is empty or still being written.")
        return

    st.caption(
        f"📁 {run_dir.name}  ·  {len(art['pptx']) + len(art['pdf'])} deck · "
        f"{len(art['excel'])} Excel · {len(art['csv'])} CSV · {len(art['charts'])} charts")

    _download_strip(art)
    st.divider()

    tabs = st.tabs(["Segment summary", "Deck", "Excel reports",
                    "Segment data", "Charts"])

    with tabs[0]:
        _render_summary(run_dir, art)

    with tabs[1]:
        if not art["pptx"] and not art["pdf"]:
            ui.callout("empty", "No deck in this run", "")
        for p in art["pptx"]:
            _render_pptx(p)
        for p in art["pdf"]:
            _render_pdf(p)

    with tabs[2]:
        if not art["excel"]:
            ui.callout("empty", "No Excel workbooks in this run", "")
        for p in art["excel"]:
            with st.expander(p.name, expanded=len(art["excel"]) == 1):
                _render_excel(p)

    with tabs[3]:
        if not art["csv"]:
            ui.callout("empty", "No CSV files in this run", "")
        for p in art["csv"]:
            with st.expander(p.name, expanded=False):
                _render_csv(p)

    with tabs[4]:
        if not art["charts"]:
            ui.callout("empty", "No charts in this run", "")
        cols = st.columns(2)
        for k, img in enumerate(art["charts"]):
            with cols[k % 2]:
                st.image(str(img), caption=img.name, width="stretch")


def _render_summary(run_dir: Path, art: dict) -> None:
    combined = next((p for p in art["csv"] if p.name.startswith("all_segments")), None)
    if combined is None:
        ui.callout("empty", "No summary available",
                   "No all_segments_*.csv in this run.")
        return
    try:
        df = pd.read_csv(combined)
    except Exception as exc:  # noqa: BLE001
        ui.callout("error", f"Could not read {combined.name}", str(exc))
        return

    n = len(df)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Outlets", f"{n:,}")
    if "cluster" in df.columns:
        c2.metric("Segments", int(df["cluster"].nunique()))
    if "priority" in df.columns:
        c3.metric("Priority tiers", int(df["priority"].nunique()))
    if "VPO" in df.columns:
        c4.metric("Total monthly VPO",
                  f"₹{pd.to_numeric(df['VPO'], errors='coerce').sum():,.0f}")

    label_col = next((c for c in ("segment_label", "cluster_label", "label")
                      if c in df.columns), None)

    # ── per-segment table ──
    if "cluster" in df.columns:
        st.markdown("###### Segments")
        rows = []
        for cid, g in df.groupby("cluster"):
            rows.append({
                "Seg": int(cid),
                "Label": (g[label_col].iloc[0] if label_col else f"Segment {cid}"),
                "Outlets": len(g),
                "% Univ": 100 * len(g) / n,
                "Avg VPO": pd.to_numeric(g.get("VPO"), errors="coerce").mean(),
                "Avg SKUs": pd.to_numeric(g.get("AVG_SKU"), errors="coerce").mean(),
            })
        seg_df = pd.DataFrame(rows).sort_values("Outlets", ascending=False)
        st.dataframe(
            seg_df, hide_index=True, width="stretch",
            column_config={
                "% Univ": st.column_config.NumberColumn(format="%.1f%%"),
                "Avg VPO": st.column_config.NumberColumn(format="₹%d"),
                "Avg SKUs": st.column_config.NumberColumn(format="%.1f"),
            },
        )
        counts = seg_df.set_index("Label")["Outlets"]
        st.bar_chart(counts, height=240)

    # ── priority table ──
    if "priority" in df.columns:
        st.markdown("###### Priority tiers")
        prows = []
        for tier in _PRIORITY_ORDER:
            g = df[df["priority"] == tier]
            if not len(g):
                continue
            prows.append({
                "Tier": tier,
                "Outlets": len(g),
                "% Univ": 100 * len(g) / n,
                "Avg actual VPO": pd.to_numeric(g.get("VPO"), errors="coerce").mean(),
                "Avg potential VPO": pd.to_numeric(
                    g.get("predicted_potential_vpo"), errors="coerce").mean(),
                "Avg gap": pd.to_numeric(g.get("opportunity_gap"),
                                         errors="coerce").mean(),
            })
        if prows:
            pdf_ = pd.DataFrame(prows)
            st.dataframe(
                pdf_, hide_index=True, width="stretch",
                column_config={
                    "% Univ": st.column_config.NumberColumn(format="%.1f%%"),
                    "Avg actual VPO": st.column_config.NumberColumn(format="₹%d"),
                    "Avg potential VPO": st.column_config.NumberColumn(format="₹%d"),
                    "Avg gap": st.column_config.NumberColumn(format="₹%d"),
                },
            )

    with st.expander("Preview combined outlet table"):
        st.dataframe(_arrow_safe(df.head(2000)), width="stretch", height=420,
                     hide_index=True)
