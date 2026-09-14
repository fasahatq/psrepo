"""Discover pipeline deliverables in ``outputs/<timestamp>/`` and the input files
in ``inbox/`` — as plain dicts the API can serialise.

Ported from ``gui/outputs.py`` (the Streamlit renderers) and ``gui/app.py``
(``_scan_inbox``), with all Streamlit widgets removed.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import pandas as pd

PROJECT_ROOT = Path(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
OUTPUTS_DIR = PROJECT_ROOT / "outputs"
INBOX_DIR = PROJECT_ROOT / "inbox"

_PRIORITY_ORDER = ["A+", "A", "B+", "B", "C+", "C", "D+", "D"]

# ── inbox classification (from gui/app.py) ───────────────────────────────────
_OUTLET_COLS = {"OUTLET_UID_EDITED", "VPO", "TOTAL_REVENUE"}
_SKU_COLS = {"CUST_UNIQ_ID_VAL", "NET_SALES", "BRND_NM"}
_PREFERRED = "Market_Master_File.csv"

_MIME = {
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".png": "image/png",
}


def fmt_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}"
        n /= 1024
    return f"{n:.0f} TB"


def mime_for(name: str) -> str:
    return _MIME.get(Path(name).suffix.lower(), "application/octet-stream")


# ── inbox ────────────────────────────────────────────────────────────────────
def scan_inbox() -> list[dict]:
    if not INBOX_DIR.is_dir():
        return []
    entries: list[dict] = []
    for p in sorted(INBOX_DIR.glob("*")):
        if p.suffix.lower() not in (".csv", ".xlsx", ".xls"):
            continue
        try:
            if p.suffix.lower() == ".csv":
                cols = set(pd.read_csv(p, nrows=0).columns.str.strip())
            else:
                cols = set(pd.read_excel(p, nrows=0).columns.str.strip())
        except Exception:
            cols = set()
        is_dataset = bool(cols & _OUTLET_COLS) or _SKU_COLS.issubset(cols)
        entries.append({
            "name": p.name,
            "size": p.stat().st_size,
            "size_h": fmt_size(p.stat().st_size),
            "kind": "dataset" if is_dataset else "reference",
        })
    entries.sort(key=lambda d: (d["kind"] != "dataset",
                                d["name"] != _PREFERRED, -d["size"]))
    return entries


# ── runs ─────────────────────────────────────────────────────────────────────
def _run_dir(run_id: str) -> Path:
    # run_id is the folder name, e.g. "20260904_045700". Guard against traversal.
    safe = os.path.basename(run_id)
    d = OUTPUTS_DIR / safe
    if not d.is_dir() or not safe[:1].isdigit():
        raise FileNotFoundError(run_id)
    return d


def run_label(name: str) -> str:
    try:
        return datetime.strptime(name, "%Y%m%d_%H%M%S").strftime("%d %b %Y  %H:%M")
    except ValueError:
        return name


def _run_ts(name: str) -> str:
    try:
        return datetime.strptime(name, "%Y%m%d_%H%M%S").isoformat()
    except ValueError:
        return ""


def scan_artifacts(run_dir: Path) -> dict:
    return {
        "pptx": sorted(run_dir.glob("*.pptx")),
        "pdf": sorted(run_dir.glob("*.pdf")),
        "excel": sorted(run_dir.glob("*.xlsx")),
        "csv": sorted(run_dir.glob("*.csv")),
        "charts": sorted((run_dir / "charts").glob("*.png"))
        if (run_dir / "charts").is_dir() else [],
    }


def _slide_count(path: Path) -> int | None:
    try:
        from pptx import Presentation
        return len(list(Presentation(str(path)).slides))
    except Exception:
        return None


def _deck_titles(path: Path) -> list[str]:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
    except Exception:
        return []
    titles: list[str] = []
    for i, sl in enumerate(prs.slides, 1):
        t = ""
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().splitlines()[0]
                break
        titles.append(f"{i}. {t or '(untitled)'}")
    return titles


def list_runs() -> list[Path]:
    if not OUTPUTS_DIR.is_dir():
        return []
    runs = [p for p in OUTPUTS_DIR.iterdir() if p.is_dir() and p.name[:1].isdigit()]
    return sorted(runs, reverse=True)


def _summary_frame(run_dir: Path):
    combined = next((p for p in run_dir.glob("all_segments*.csv")), None)
    if combined is None:
        return None
    try:
        return pd.read_csv(combined)
    except Exception:
        return None


def run_summary(run_dir: Path) -> dict:
    """Headline numbers + per-segment / per-priority tables from all_segments*.csv."""
    df = _summary_frame(run_dir)
    if df is None:
        return {}
    n = len(df)
    out: dict = {"outlets": int(n)}

    if "cluster" in df.columns:
        out["segment_count"] = int(df["cluster"].nunique())
    if "priority" in df.columns:
        out["priority_tiers"] = int(df["priority"].nunique())
    if "VPO" in df.columns:
        out["total_vpo"] = float(pd.to_numeric(df["VPO"], errors="coerce").sum())

    label_col = next((c for c in ("segment_label", "cluster_label", "label")
                      if c in df.columns), None)

    segments: list[dict] = []
    if "cluster" in df.columns:
        for cid, g in df.groupby("cluster"):
            segments.append({
                "id": int(cid),
                "label": str(g[label_col].iloc[0]) if label_col else f"Segment {cid}",
                "outlets": int(len(g)),
                "pct_universe": round(100 * len(g) / n, 1),
                "avg_vpo": _safe_mean(g.get("VPO")),
                "avg_skus": _safe_mean(g.get("AVG_SKU")),
                "avg_gap": _safe_mean(g.get("opportunity_gap")),
                "description": (str(g["segment_description"].iloc[0])
                               if "segment_description" in g.columns else ""),
                "action": (str(g["segment_action"].iloc[0])
                           if "segment_action" in g.columns else ""),
            })
        segments.sort(key=lambda s: s["outlets"], reverse=True)
    out["segments"] = segments

    tiers: list[dict] = []
    if "priority" in df.columns:
        for tier in _PRIORITY_ORDER:
            g = df[df["priority"] == tier]
            if not len(g):
                continue
            tiers.append({
                "tier": tier,
                "outlets": int(len(g)),
                "pct_universe": round(100 * len(g) / n, 1),
                "avg_actual_vpo": _safe_mean(g.get("VPO")),
                "avg_potential_vpo": _safe_mean(g.get("predicted_potential_vpo")),
                "avg_gap": _safe_mean(g.get("opportunity_gap")),
            })
    out["priority_tiers_table"] = tiers
    return out


def _safe_mean(series) -> float | None:
    if series is None:
        return None
    v = pd.to_numeric(series, errors="coerce").mean()
    return None if pd.isna(v) else float(v)


def run_brief(run_dir: Path) -> dict:
    """Lightweight entry for the runs list."""
    art = scan_artifacts(run_dir)
    name = run_dir.name
    summ = run_summary(run_dir)
    deck = art["pptx"][0] if art["pptx"] else None
    return {
        "id": name,
        "label": run_label(name),
        "ts": _run_ts(name),
        "market": "India",  # no market field in the data model — see plan
        "counts": {
            "deck": len(art["pptx"]) + len(art["pdf"]),
            "excel": len(art["excel"]),
            "csv": len(art["csv"]),
            "charts": len(art["charts"]),
        },
        "deck_name": deck.name if deck else None,
        "deck_slides": _slide_count(deck) if deck else None,
        "outlets": summ.get("outlets"),
        "segment_count": summ.get("segment_count"),
        "total_vpo": summ.get("total_vpo"),
        "status": "Ready",
    }


_ASSET_WORDS = ("chiller", "cooler", "shelf", "planogram", "posm", "display",
                "rack", "counter", "merch", "space", "fridge", "visi")


def _split_actions(raw: str) -> list[dict]:
    """'1. [Portfolio] Do a thing.  (KPI: x)\n2. ...' -> structured items."""
    import re
    out: list[dict] = []
    if not isinstance(raw, str):
        return out
    for chunk in re.split(r"\s*\d+\.\s+", raw.strip()):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.match(r"\[([^\]]+)\]\s*(.+)", chunk) or re.match(r"([A-Za-z][\w &/]{1,28}?):\s+(.+)", chunk)
        lever = (m.group(1).strip() if m else "")
        body = (m.group(2).strip() if m else chunk)
        kpi = ""
        km = re.search(r"\(KPI:\s*(.+?)\)\s*$", body)
        if km:
            kpi = km.group(1).strip()
            body = body[: km.start()].strip()
        blob = f"{lever} {body}".lower()
        out.append({
            "lever": lever, "text": body, "kpi": kpi,
            "is_asset": any(w in blob for w in _ASSET_WORDS),
        })
    return out


def segment_cards(run_dir: Path) -> list[dict]:
    """The 'Segment Cards' sheet of segment_report*.xlsx — the crisp per-segment
    view: hero SKUs, top actions, and the merch/space assets called out."""
    rep = next((p for p in run_dir.glob("segment_report*.xlsx")), None)
    if rep is None:
        return []
    try:
        df = pd.read_excel(rep, sheet_name="Segment Cards")
    except Exception:
        return []
    cards: list[dict] = []
    for _, r in df.iterrows():
        actions = _split_actions(r.get("Top Actions", ""))
        skus = [s.strip() for s in str(r.get("Hero SKUs", "")).split(",") if s.strip()]
        cards.append({
            "cluster": int(r["Cluster"]) if pd.notna(r.get("Cluster")) else None,
            "label": str(r.get("Label", "")),
            "channel": str(r.get("Channel", "")),
            "occasion": str(r.get("Occasion", "")),
            "headline": str(r.get("Headline", "")),
            "snapshot": str(r.get("Shopper Snapshot", "")),
            "growth": str(r.get("Growth", "")),
            "dominant_sec": str(r.get("Dominant SEC", "")),
            "hero_skus": skus,
            "actions": actions,
            "assets": [a for a in actions if a["is_asset"]],
        })
    return cards


def top_skus_across_segments(cards: list[dict], limit: int = 12) -> list[dict]:
    """Hero SKUs ranked by how many segments call them out."""
    from collections import Counter, defaultdict
    seen: dict[str, list[str]] = defaultdict(list)
    for c in cards:
        for s in c["hero_skus"]:
            seen[s].append(c["label"])
    counts = Counter({k: len(v) for k, v in seen.items()})
    return [{"sku": k, "segments": seen[k], "count": n}
            for k, n in counts.most_common(limit)]


INR_PER_USD = 83.0  # documented fixed rate; revisit if a live FX rate is ever wired in


def _with_financials(cards: list[dict], summary: dict) -> list[dict]:
    """Attach each card's outlet count / opportunity gap (from the summary's per-
    segment table) plus the derived USD upside figures the Output Studio cards show."""
    by_cluster = {s["id"]: s for s in summary.get("segments", [])}
    for card in cards:
        s = by_cluster.get(card.get("cluster")) or {}
        outlets = s.get("outlets")
        avg_gap = s.get("avg_gap")
        # Upside is floored at 0 — a negative gap means these outlets already exceed
        # their modelled potential (agents/prioritization_agent.py's `has_upside`
        # convention), which isn't a "revenue opportunity" to report on a card.
        upside_monthly = (max(avg_gap, 0) * outlets / INR_PER_USD
                          if avg_gap is not None and outlets else None)
        card["outlets"] = outlets
        card["avg_gap_inr"] = avg_gap
        card["upside_usd_monthly"] = upside_monthly
        card["revenue_impact_usd_annual"] = (upside_monthly * 12
                                             if upside_monthly is not None else None)
    return cards


def run_detail(run_id: str) -> dict:
    run_dir = _run_dir(run_id)
    art = scan_artifacts(run_dir)
    files: list[dict] = []
    for key in ("pptx", "pdf", "excel", "csv"):
        for p in art[key]:
            entry = {
                "name": p.name,
                "kind": key,
                "size": p.stat().st_size,
                "size_h": fmt_size(p.stat().st_size),
            }
            if key == "pptx":
                entry["slides"] = _slide_count(p)
            files.append(entry)
    deck = art["pptx"][0] if art["pptx"] else None
    summary = run_summary(run_dir)
    cards = _with_financials(segment_cards(run_dir), summary)
    return {
        **run_brief(run_dir),
        "files": files,
        "charts": [c.name for c in art["charts"]],
        "deck_titles": _deck_titles(deck) if deck else [],
        "summary": summary,
        "segment_cards": cards,
        "top_skus": top_skus_across_segments(cards),
    }


def artifact_path(run_id: str, name: str) -> Path:
    run_dir = _run_dir(run_id)
    safe = os.path.basename(name)
    p = run_dir / safe
    if not p.is_file():
        raise FileNotFoundError(name)
    return p


def chart_path(run_id: str, name: str) -> Path:
    run_dir = _run_dir(run_id)
    safe = os.path.basename(name)
    p = run_dir / "charts" / safe
    if not p.is_file():
        raise FileNotFoundError(name)
    return p
